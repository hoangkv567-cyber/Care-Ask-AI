import os, re, sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

md_path = r"C:\Users\hoang\.gemini\antigravity\brain\3b8a2a5d-bbc8-4cdf-9e87-067a4f33fb51\bao_cao_tien_do_tcm_system.md"
docx_paths = [
    r"C:\Users\hoang\Downloads\VuNguyenQuocHoang_N3.docx",
    r"C:\Users\hoang\TTDN\TCM_System\Bao_Cao_Tien_Do_TCM_System.docx"
]

with open(md_path, "r", encoding="utf-8") as f:
    text = f.read()

def build_docx(target_path):
    doc = docx.Document()

    # Set Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Normal Style Font
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Arial'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def parse_inline(p, line):
        parts = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('`') and part.endswith('`'):
                run = p.add_run(part[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
            elif part.startswith('*') and part.endswith('*'):
                run = p.add_run(part[1:-1])
                run.italic = True
            else:
                p.add_run(part)

    def render_table(table_lines):
        if not table_lines:
            return
        rows = []
        for line in table_lines:
            if re.match(r'^\s*\|?\s*:?-+:?\s*\|', line):
                continue
            cells = [c.strip() for c in line.strip('|').split('|')]
            if cells:
                rows.append(cells)
        if not rows:
            return
        
        table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, row_data in enumerate(rows):
            for c_idx, cell_value in enumerate(row_data):
                cell = table.cell(r_idx, c_idx)
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                parse_inline(p, cell_value)
                
                if r_idx == 0:
                    set_cell_background(cell, "0F766E")
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.bold = True
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    else:
                        set_cell_background(cell, "FFFFFF")

    def add_image_figure(img_path, caption_text, width_inches=6.0):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run()
            run.add_picture(img_path, width=Inches(width_inches))
            
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_before = Pt(2)
            cap.paragraph_format.space_after = Pt(12)
            r_cap = cap.add_run(caption_text)
            r_cap.font.name = 'Arial'
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    lines = text.split('\n')
    i = 0
    in_code_block = False
    code_lines = []
    table_lines = []
    toc_inserted = False

    while i < len(lines):
        line = lines[i]
        
        # Check image tags in markdown
        if line.strip().startswith("![Hình 1"):
            add_image_figure("images/diagram1_kg_schema.png", "Hình 1: Sơ đồ Mối quan hệ giữa các Thực thể trong Đồ thị Tri thức Đông y", width_inches=4.8)
            i += 1
            continue
        elif line.strip().startswith("![Hình 2"):
            add_image_figure("images/diagram2_graph_rag_flow.png", "Hình 2: Sơ đồ Luồng Suy luận và Truy xuất Ngữ cảnh Graph RAG", width_inches=6.2)
            i += 1
            continue
        elif line.strip().startswith("![Hình 3"):
            add_image_figure("images/diagram3_vlm_fallback.png", "Hình 3: Sơ đồ Chuỗi Fallback VLM 3 Tầng Tự động", width_inches=5.2)
            i += 1
            continue
        elif line.strip().startswith("![Hình 4"):
            add_image_figure("images/diagram4_system_architecture.png", "Hình 4: Sơ đồ Kiến trúc Tổng thể Hệ thống Microservices & AI Engine", width_inches=5.5)
            i += 1
            continue

        if line.startswith("```"):
            if in_code_block:
                code_text = "\n".join(code_lines)
                # Code snippet rendering
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
        
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
        
        # Table detection
        if '|' in line and not line.strip().startswith('#'):
            table_lines.append(line)
            i += 1
            continue
        else:
            if table_lines:
                render_table(table_lines)
                table_lines = []
        
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        
        # Insert Table of Contents before Section I
        if "## I. CƠ SỞ LÝ THUYẾT" in stripped and not toc_inserted:
            toc_inserted = True
            
            # Header MỤC LỤC
            p_th = doc.add_paragraph()
            p_th.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_th.paragraph_format.space_before = Pt(14)
            p_th.paragraph_format.space_after = Pt(10)
            r_th = p_th.add_run("MỤC LỤC")
            r_th.font.name = 'Arial'
            r_th.font.size = Pt(16)
            r_th.font.bold = True
            r_th.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
            
            # Native XML TOC Field
            p_xml_toc = doc.add_paragraph()
            fldSimple = parse_xml(r'<w:fldSimple %s w:instr="TOC \o &quot;1-3&quot; \h \z \u"/>' % nsdecls('w'))
            p_xml_toc._p.append(fldSimple)

            toc_entries = [
                ("I. CƠ SỞ LÝ THUYẾT & NỀN TẢNG Y KHOA — CÔNG NGHỆ AI", "1", 1, True),
                ("1. Lý luận Y học Cổ truyền (Đông y) và Phương pháp Tứ Chẩn Số", "1", 2, False),
                ("2. Đồ thị Tri thức (Knowledge Graph) và Kiến trúc Graph RAG", "2", 2, False),
                ("2.1. Cấu trúc Thực thể (Nodes) trong Đồ thị Tri thức Đông y (TCM-KG)", "2", 3, False),
                ("2.2. Mạng lưới Mối quan hệ (Edges) trong Đồ thị Tri thức", "2", 3, False),
                ("2.3. Quy trình Suy luận Biện chứng Luận trị trên Đồ thị", "3", 3, False),
                ("3. Kiến trúc Graph RAG (Retrieval-Augmented Generation kết hợp Đồ thị)", "3", 2, False),
                ("4. Kiến trúc Đa phương thức & Hệ sinh thái Ứng dụng Di động", "4", 2, False),
                ("4.1. Chuỗi Fallback VLM Đa tầng (Multi-tier VLM Fallback Chain)", "4", 3, False),
                ("4.2. Ứng dụng Di động React Native CLI (DongY Native App)", "4", 3, False),
                ("II. CẤU TRÚC KIẾN TRÚC HỆ THỐNG VÀ THƯ VIỆN CÔNG NGHỆ", "5", 1, True),
                ("III. QUÁ TRÌNH PHÁT TRIỂN & TIẾN ĐỘ THEO TUẦN (TUẦN 1 — TUẦN 8)", "6", 1, True),
                ("Tuần 1: Nghiên cứu TCM Knowledge Graph & Kiến trúc Graph RAG", "6", 2, False),
                ("Tuần 2: Xây dựng & Nạp dữ liệu vào TCM Knowledge Graph (Neo4j)", "6", 2, False),
                ("Tuần 3: Phát triển Module Xử lý Đa phương thức (Visual Input)", "6", 2, False),
                ("Tuần 4: Xây dựng Hệ thống Hỏi đáp Đồ thị (Text-to-Cypher & Graph QA)", "7", 2, False),
                ("Tuần 5: Sinh Câu trả lời có Giải thích (Explainable AI Generation)", "7", 2, False),
                ("Tuần 6: Xây dựng Ứng dụng Di động React Native & Web UI", "7", 2, False),
                ("Tuần 7: Kiểm thử Lâm sàng & Tinh chỉnh Logic Đông Y Chuyên sâu", "7", 2, False),
                ("Tuần 8: Đóng gói Mã nguồn, Tối ưu hóa & Báo cáo Tổng kết", "8", 2, False),
                ("IV. CƠ SỞ THỰC HÀNH VÀ KẾT QUẢ THỰC NGHIỆM LÂM SÀNG", "8", 1, True),
                ("1. Mã Nguồn Xử lý Ứng dụng Di động React Native (DongY/App.tsx)", "8", 2, False),
                ("2. Thực thi Truy vấn Đồ thị Tri thức Neo4j (Cypher Code)", "9", 2, False),
                ("3. Kết quả Chẩn đoán Lâm sàng Mẫu từ Hệ thống (System Diagnostic)", "9", 2, False),
                ("V. CÁC VẤN ĐỀ KỸ THUẬT ĐÃ GIẢI QUYẾT VÀ ĐÁNH GIÁ TỔNG KẾT", "10", 1, True),
                ("KẾT LUẬN TỔNG THỂ", "11", 2, False),
            ]

            figure_entries = [
                ("Hình 1: Sơ đồ Mối quan hệ giữa các Thực thể trong Đồ thị Tri thức", "2"),
                ("Hình 2: Sơ đồ Luồng Suy luận và Truy xuất Ngữ cảnh Graph RAG", "3"),
                ("Hình 3: Sơ đồ Chuỗi Fallback VLM 3 Tầng Tự động", "4"),
                ("Hình 4: Sơ đồ Kiến trúc Tổng thể Hệ thống Microservices & AI Engine", "5"),
            ]

            for title, page, level, is_bold in toc_entries:
                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_before = Pt(3 if level == 1 else 1)
                p_item.paragraph_format.space_after = Pt(2)
                p_item.paragraph_format.left_indent = Inches(0.0 if level == 1 else (0.25 if level == 2 else 0.5))
                
                r_t = p_item.add_run(title + " ")
                r_t.font.name = 'Arial'
                r_t.font.size = Pt(10.5 if level == 1 else 10)
                r_t.font.bold = is_bold
                r_t.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) if level == 1 else RGBColor(0x33, 0x41, 0x55)
                
                dots_count = max(5, 85 - len(title) - (level * 4))
                r_d = p_item.add_run(" " + ". " * (dots_count // 2) + " ")
                r_d.font.name = 'Arial'
                r_d.font.size = Pt(8.5)
                r_d.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
                
                r_p = p_item.add_run(page)
                r_p.font.name = 'Arial'
                r_p.font.size = Pt(10)
                r_p.font.bold = is_bold
                r_p.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) if level == 1 else RGBColor(0x47, 0x55, 0x69)

            p_fh = doc.add_paragraph()
            p_fh.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_fh.paragraph_format.space_before = Pt(14)
            p_fh.paragraph_format.space_after = Pt(8)
            r_fh = p_fh.add_run("DANH MỤC HÌNH VẼ & SƠ ĐỒ")
            r_fh.font.name = 'Arial'
            r_fh.font.size = Pt(14)
            r_fh.font.bold = True
            r_fh.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

            for title, page in figure_entries:
                p_fig = doc.add_paragraph()
                p_fig.paragraph_format.space_before = Pt(2)
                p_fig.paragraph_format.space_after = Pt(2)
                p_fig.paragraph_format.left_indent = Inches(0.15)
                
                r_t = p_fig.add_run(title + " ")
                r_t.font.name = 'Arial'
                r_t.font.size = Pt(10)
                r_t.font.italic = True
                r_t.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                
                dots_count = max(5, 80 - len(title))
                r_d = p_fig.add_run(" " + ". " * (dots_count // 2) + " ")
                r_d.font.name = 'Arial'
                r_d.font.size = Pt(8.5)
                r_d.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
                
                r_p = p_fig.add_run(page)
                r_p.font.name = 'Arial'
                r_p.font.size = Pt(10)
                r_p.font.bold = True
                r_p.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

            div = doc.add_paragraph()
            div.paragraph_format.space_before = Pt(8)
            div.paragraph_format.space_after = Pt(12)
            run_div = div.add_run("_________________________________________________________________________________")
            run_div.font.color.rgb = RGBColor(0xCC, 0xFB, 0xF1)

        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(stripped[2:])
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.style = doc.styles['Heading 1']
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[3:])
            run.bold = True
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0x11, 0x5E, 0x59)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.style = doc.styles['Heading 2']
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[4:])
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.style = doc.styles['Heading 3']
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[5:])
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0x14, 0xB8, 0xA6)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            parse_inline(p, stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            content = re.sub(r'^\d+\.\s', '', stripped)
            parse_inline(p, content)
        elif stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("_________________________________________________________________________________")
            run.font.color.rgb = RGBColor(0xCC, 0xFB, 0xF1)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            parse_inline(p, stripped)
        
        i += 1

    if table_lines:
        render_table(table_lines)

    doc.save(target_path)
    print("SUCCESSFULLY REBUILT CLEAN DOCX AT:", target_path)

for path in docx_paths:
    try:
        build_docx(path)
    except Exception as e:
        print(f"Error building {path}: {e}")
